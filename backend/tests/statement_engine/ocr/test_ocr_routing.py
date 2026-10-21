"""T4 routing through the REAL orchestrator: image-only PDF + fixture OCR service (no model, no GPU)."""
import json
import os

import pytest

from app.services.pipeline_orchestrator import PipelineAbortError, process_statement
from app.services.statement_engine.ocr.client import _sha256
from app.services.statement_engine.ocr.contracts import document_to_dict
from tests.statement_engine.ocr.synthetic import make_document


def _scan_pdf(path, pages=2, tag=b"a"):
    from PIL import Image
    imgs = [Image.new("RGB", (300 + len(tag), 400), "white") for _ in range(pages)]
    imgs[0].save(path, "PDF", save_all=True, append_images=imgs[1:], resolution=100)
    return str(path)


def _env(monkeypatch, tmp_path, mode="serve", url=True):
    monkeypatch.setenv("STATEMENT_ENGINE_OCR_MODE", mode)
    if url:
        monkeypatch.setenv("STATEMENT_ENGINE_OCR_URL", f"mock:{tmp_path}")
    else:
        monkeypatch.delenv("STATEMENT_ENGINE_OCR_URL", raising=False)
    monkeypatch.setenv("STATEMENT_ENGINE_UNIVERSAL_MODE", "off")
    monkeypatch.setenv("STATEMENT_ENGINE_V2_MODE", "off")


def _fixture(tmp_path, pdf, doc):
    (tmp_path / f"{_sha256(pdf)}.json").write_text(json.dumps(document_to_dict(doc)))


def _run(pdf, tmp_path):
    out = tmp_path / "out"; out.mkdir(exist_ok=True)
    return process_statement(pdf, {"full_name": "T", "account_type": "Salaried", "bank_name": "Other bank"},
                             mode="free", output_dir=str(out))


def test_flag_off_scanned_file_behaves_as_before(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path, mode="off")
    pdf = _scan_pdf(tmp_path / "s.pdf")
    with pytest.raises(Exception) as ei:
        _run(pdf, tmp_path)
    assert "ocr" not in str(ei.value).lower() or "not" in str(ei.value).lower()
    assert getattr(ei.value, "error_code", "") != "OCR_FAILED"


def test_serve_scanned_statement_is_processed_and_capped_at_review(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    pdf = _scan_pdf(tmp_path / "s.pdf")
    doc, rows = make_document(n_rows=24, per_page=12)
    _fixture(tmp_path, pdf, doc)
    res = _run(pdf, tmp_path)
    assert res["status"] == "success"
    se = res["statement_engine"]
    assert se["engine"] == "ocr" and se["verdict"] == "REVIEW" and se["needs_review"] is True and se["anchored"] is False
    assert "OCR_READER" in se["review_codes"]
    assert se["ocr"]["pages_ok"] == [1, 2] and se["ocr"]["failed_pages"] == {} and se["ocr"]["partial"] is False
    assert res["stats"]["total_transactions"] == len(rows) == 24 and res["stats"]["coverage_percent"] == 100.0
    import glob, openpyxl
    wb = openpyxl.load_workbook(res["excel_path"])
    assert "Transactions" in wb.sheetnames and "Statement Check" in wb.sheetnames
    cells = {r[0]: r[1] for r in wb["Statement Check"].iter_rows(values_only=True) if r and r[0]}
    assert cells["Engine"] == "ocr" and cells["Independently verified"] == "no" and cells["Pages NOT read"] == "none"


def test_partial_pages_are_reported_and_never_look_complete(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    pdf = _scan_pdf(tmp_path / "s.pdf", pages=3)
    doc, _ = make_document(n_rows=36, per_page=12)
    doc.pages[2].status, doc.pages[2].error, doc.pages[2].blocks = "failed", "OCR_TIMEOUT", []
    _fixture(tmp_path, pdf, doc)
    try:
        res = _run(pdf, tmp_path)
    except PipelineAbortError as e:
        assert e.stage == "ocr_reader"           # the existing checks refused a gapped statement: acceptable
        return
    se = res["statement_engine"]
    assert se["verdict"] == "REVIEW" and se["ocr"]["partial"] is True
    assert se["ocr"]["failed_pages"] == {"3": "OCR_TIMEOUT"} and "OCR_PARTIAL_PAGES" in se["review_codes"]


def test_serve_with_unavailable_service_fails_honestly_with_an_ocr_code(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    pdf = _scan_pdf(tmp_path / "s.pdf")          # no fixture -> OCR_NO_FIXTURE
    with pytest.raises(PipelineAbortError) as ei:
        _run(pdf, tmp_path)
    assert ei.value.stage == "ocr_reader" and ei.value.error_code == "OCR_NO_FIXTURE"


def test_shadow_mode_records_the_block_and_serves_nothing(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path, mode="shadow")
    pdf = _scan_pdf(tmp_path / "s.pdf")
    doc, _ = make_document()
    _fixture(tmp_path, pdf, doc)
    with pytest.raises(Exception) as ei:          # still fails exactly as today
        _run(pdf, tmp_path)
    assert getattr(ei.value, "stage", "") != "ocr_reader"


def test_text_pdf_never_touches_ocr(monkeypatch, tmp_path):
    """A text-layer PDF must not reach the OCR client, or even the capability check, with the flag on."""
    _env(monkeypatch, tmp_path)
    import app.services.statement_engine.ocr.client as c
    called = []
    monkeypatch.setattr(c, "ocr_pdf", lambda *a, **k: called.append(1))
    from reportlab.pdfgen import canvas
    pdf = str(tmp_path / "t.pdf")
    cv = canvas.Canvas(pdf)
    for i in range(60):
        cv.drawString(40, 800 - i * 12, f"01/04/2026 UPI payment {i} 1,000.00 24,000.00")
    cv.save()
    try:
        _run(pdf, tmp_path)
    except Exception:
        pass                                      # the outcome for this toy file is irrelevant: OCR must not run
    assert called == []


def test_partial_statement_check_sheet_lists_the_failed_pages(tmp_path):
    from app.services.statement_engine.serve import append_statement_check_sheet
    import openpyxl
    f = str(tmp_path / "w.xlsx"); openpyxl.Workbook().save(f)
    block = {"engine": "ocr", "verdict": "REVIEW", "needs_review": True, "anchored": False,
             "review_codes": ["OCR_READER", "OCR_PARTIAL_PAGES"],
             "ocr": {"ocr_engine": "PaddleOCR-VL-1.6", "pages_ok": [1, 3], "failed_pages": {"2": "OCR_TIMEOUT"},
                     "verification": {"status": "ran", "amounts_checked": 4, "confirmed": 4}}}
    assert append_statement_check_sheet(f, block)
    cells = {r[0]: r[1] for r in openpyxl.load_workbook(f)["Statement Check"].iter_rows(values_only=True) if r and r[0]}
    notes = " ".join(str(r[1]) for r in openpyxl.load_workbook(f)["Statement Check"].iter_rows(values_only=True) if r and r[0] == "Note")
    assert cells["Pages NOT read"] == "2 (OCR_TIMEOUT)" and "does NOT contain every transaction" in notes
