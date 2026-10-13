"""V2 serving: policy matrix (contract §2), review codes (§4), block shape (§3),
allowlist/mode gating (§1), bank mismatch (§6), orchestrator integration, and the
Excel Statement Check sheet. Synthetic objects only — no PDFs, no PII."""
import os
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.services.statement_engine import serve, shadow
from app.services.statement_engine.v2_router import RoutedExtraction


# --------------------------------------------------------------------------- #
# Fixtures / fakes                                                             #
# --------------------------------------------------------------------------- #

@pytest.fixture(autouse=True)
def _clean_env(monkeypatch):
    monkeypatch.delenv(shadow.MODE_ENV, raising=False)
    monkeypatch.delenv(serve.SERVE_BANKS_ENV, raising=False)


def _checks(**statuses):
    return {name: {"status": st} for name, st in statuses.items()}


def _routed(verdict="VERIFIED", layout_id="hdfc_x", status="v2", warnings=(),
            checks=None, txns=3, bank_key="hdfc"):
    extras = {
        "verdict": verdict,
        "verdict_reasons": [],
        "checks": checks if checks is not None else _checks(
            schema="pass", chain="pass", side="pass", reference="pass",
            totals="pass", dates="pass", period="pass", date_outliers="pass",
        ),
        "printed_summary": {"opening": "1000.00", "closing": "900.00",
                            "debits": "500.00", "credits": "400.00"},
        "ifsc": "HDFC0000001",
    }
    meta = SimpleNamespace(
        bank_key=bank_key, opening_balance=Decimal("1000"),
        closing_balance=Decimal("900"), statement_from="2023-04-01",
        statement_to="2023-04-30", account_number="9999", extras=extras,
    )
    tx = [
        SimpleNamespace(to_dict=lambda i=i: {
            "date": "2023-04-0%d" % (i + 1), "description": "T%d" % i,
            "debit": 100.0 if i % 2 == 0 else None,
            "credit": None if i % 2 == 0 else 100.0,
            "balance": 900.0 + i, "ref_no": "R%d" % i, "category": "Others",
        })
        for i in range(txns)
    ]
    res = SimpleNamespace(metadata=meta, transactions=tx, warnings=list(warnings))
    return RoutedExtraction(status=status, layout_id=layout_id, result=res)


def _det(bank_id="hdfc", confidence="VERY_HIGH", score=80):
    from app.services.statement_engine.contracts import Confidence
    return SimpleNamespace(bank_id=bank_id,
                           confidence=getattr(Confidence, confidence),
                           score=score, runner_up_id=None, runner_up_score=0,
                           evidence=["marker"])


# --------------------------------------------------------------------------- #
# Mode / allowlist gating (contract §1)                                        #
# --------------------------------------------------------------------------- #

def test_mode_off_by_default():
    assert serve.v2_mode() == "off"
    assert serve.serving_enabled_for("hdfc") is False


def test_serve_needs_mode_and_allowlist(monkeypatch):
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "")
    assert serve.serving_enabled_for("federal") is False  # empty allowlist = none
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "federal,bob")
    assert serve.serving_enabled_for("federal") is True
    assert serve.serving_enabled_for("bank_of_baroda") is True   # pipeline key -> bob
    assert serve.serving_enabled_for("hdfc") is False


def test_kill_switch_between_calls(monkeypatch):
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "federal")
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    assert serve.serving_enabled_for("federal") is True
    monkeypatch.setenv(shadow.MODE_ENV, "off")
    assert serve.serving_enabled_for("federal") is False


def test_unsupported_mode_is_off(monkeypatch):
    monkeypatch.setenv(shadow.MODE_ENV, "nonsense")
    assert serve.v2_mode() == "off"
    assert serve.serving_enabled_for("federal") is False


def test_spec_bank_keys_and_reverse_map():
    keys = serve.spec_bank_keys()
    assert {"hdfc", "sbi", "federal"} <= keys
    assert serve.pipeline_key_for("bob") == "bank_of_baroda"
    assert serve.pipeline_key_for("hdfc") == "hdfc"


# --------------------------------------------------------------------------- #
# Review codes (contract §4)                                                   #
# --------------------------------------------------------------------------- #

def test_review_codes_mapping():
    r = _routed("REVIEW", checks=_checks(
        schema="fail", chain="fail", side="warn", totals="warn",
        reference="skipped", dates="warn", date_outliers="warn",
        period="warn"), warnings=["w1"])
    codes = serve.review_codes(r)
    assert codes == ["NO_INDEPENDENT_ANCHOR", "CHAIN_BREAK", "SIDE_MISMATCH",
                     "TOTALS_MISMATCH", "DATE_ORDER", "DATE_OUTLIER",
                     "OUTSIDE_PERIOD", "SCHEMA_ERRORS", "EXTRACTION_WARNINGS"]


def test_review_codes_reference_failed():
    r = _routed("FAILED", checks=_checks(reference="fail", chain="pass"))
    assert serve.review_codes(r) == ["REFERENCE_MISMATCH"]


def test_review_codes_empty_for_clean_verified():
    assert serve.review_codes(_routed("VERIFIED")) == []
    assert serve.anchored(_routed("VERIFIED")) is True


def test_anchored_false_when_reference_skipped():
    r = _routed("REVIEW", checks=_checks(reference="skipped"))
    assert serve.anchored(r) is False


# --------------------------------------------------------------------------- #
# Policy matrix (contract §2)                                                  #
# --------------------------------------------------------------------------- #

def _plan(verdict="VERIFIED", legacy=True, det=None, spec_key="hdfc", status="v2"):
    routed = _routed(verdict, status=status)
    if status != "v2":
        routed = RoutedExtraction(status="fallback", fallback_reason="no_layout")
    return serve.decide_serve(bank_key="hdfc", spec_key=spec_key, routed=routed,
                              det=det if det is not None else _det("hdfc"),
                              legacy_available=legacy, routed_ms=1.0)


def test_verified_serves_v2():
    p = _plan("VERIFIED")
    assert p.serve_v2 and p.block["engine"] == "v2"
    assert p.block["verdict"] == "VERIFIED" and p.block["anchored"] is True
    assert p.block["needs_review"] is False


def test_review_with_legacy_serves_legacy_plus_advisory():
    p = _plan("REVIEW", legacy=True)
    assert not p.serve_v2 and p.block["engine"] == "legacy"
    assert p.block["advisory"]["verdict"] == "REVIEW"
    assert p.block["verdict"] is None


def test_review_without_legacy_serves_v2_flagged():
    p = _plan("REVIEW", legacy=False)
    assert p.serve_v2 and p.needs_review
    assert p.block["engine"] == "v2" and p.block["needs_review"] is True
    assert p.block["verdict"] == "REVIEW"


def test_failed_never_serves_v2():
    p = _plan("FAILED", legacy=True)
    assert not p.serve_v2 and p.block["engine"] == "legacy"
    assert p.block["advisory"]["verdict"] == "FAILED"


def test_failed_without_legacy_keeps_error_path():
    p = _plan("FAILED", legacy=False)
    assert not p.serve_v2 and p.block["engine"] == "legacy"


def test_fallback_outcome_is_legacy_unchanged():
    p = _plan(status="fallback")
    assert not p.serve_v2 and p.block["engine"] == "legacy"
    assert p.block["advisory"] is None


def test_route_none_is_legacy_unchanged():
    p = serve.decide_serve(bank_key="hdfc", spec_key="hdfc", routed=None, det=None,
                           legacy_available=True, routed_ms=0.0)
    assert not p.serve_v2 and p.block["engine"] == "legacy"


# --------------------------------------------------------------------------- #
# Bank mismatch (contract §6)                                                  #
# --------------------------------------------------------------------------- #

def test_bank_mismatch_blocks_serving_and_marks_code():
    p = serve.decide_serve(bank_key="hdfc", spec_key="hdfc",
                           routed=_routed("VERIFIED"), det=_det("sbi"),
                           legacy_available=True, routed_ms=1.0)
    assert not p.serve_v2 and p.mismatch
    assert p.block["engine"] == "legacy"
    assert "BANK_MISMATCH" in p.block["review_codes"]
    assert p.block["detected_bank"]["source"] == "detector"
    assert p.block["detected_bank"]["key"] == "sbi"


def test_no_mismatch_when_detector_agrees():
    assert not serve.bank_mismatch(_det("hdfc"), "hdfc")
    assert not serve.bank_mismatch(_det("bank_of_baroda"), "bob")


def test_user_selection_block_source():
    b = serve.detected_bank_block(_det("hdfc"), "hdfc", mismatch=False)
    assert b["source"] == "user" and b["key"] == "hdfc"


# --------------------------------------------------------------------------- #
# Auto-detect eligibility                                                      #
# --------------------------------------------------------------------------- #

def test_autodetect_serves_only_allowlisted_routed_bank(monkeypatch):
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "federal")
    routed = _routed("VERIFIED", bank_key="federal")
    p = serve.decide_serve(bank_key="unknown", spec_key="unknown", routed=routed,
                           det=_det("federal"), legacy_available=False, routed_ms=1.0)
    assert p.serve_v2 and p.v2_spec_key == "federal"

    routed2 = _routed("VERIFIED", bank_key="hdfc")
    p2 = serve.decide_serve(bank_key="unknown", spec_key="unknown", routed=routed2,
                            det=_det("hdfc"), legacy_available=True, routed_ms=1.0)
    assert not p2.serve_v2 and p2.block["engine"] == "legacy"


# --------------------------------------------------------------------------- #
# Row source                                                                   #
# --------------------------------------------------------------------------- #

def test_row_source_from_routed():
    rs = serve.row_source_from_routed(_routed("VERIFIED"))
    assert rs is not None and len(rs.transactions) == 3
    assert rs.opening_balance == Decimal("1000.00")
    assert rs.closing_balance == Decimal("900.00")
    assert rs.account_number == "9999" and rs.ifsc == "HDFC0000001"
    assert rs.total_debits == Decimal("500.00")
    assert serve.row_source_from_routed(None) is None


# --------------------------------------------------------------------------- #
# Orchestrator integration (stubbed router + processor)                        #
# --------------------------------------------------------------------------- #

def _stub_process_ok(self, file_path, user_info, output_dir=None):
    from app.services.banks._shared.base_processor import BasePipelineResult
    from app.services.banks._shared.generic_bank import GenericProcessingMetrics
    m = GenericProcessingMetrics()
    m.transaction_count = 7
    return BasePipelineResult(status="success", excel_path=None, transactions=[{}] * 7,
                              aggregation=None, metrics=m, bank_key=self.CONFIG.bank_key)


_V2_ROWS_SEEN = []


def _stub_v2_pipeline(self, file_path, user_info, output_dir=None, row_source=None):
    _V2_ROWS_SEEN.append(row_source)
    res = _stub_process_ok(self, file_path, user_info, output_dir)
    res.metrics.transaction_count = len(row_source.transactions)
    return res


class _StubProcessor:
    CONFIG = SimpleNamespace(bank_key="hdfc")

    def __init__(self, **kw):
        pass

    def process(self, **kw):  # placeholder; tests monkeypatch it
        raise NotImplementedError

    def _run_pipeline(self, **kw):  # placeholder; tests monkeypatch it
        raise NotImplementedError


def _patch_orchestrator(monkeypatch, routed):
    import app.services.pipeline_orchestrator as po
    monkeypatch.setattr(serve, "route_for_serve", lambda p, k, timeout_s=0: (routed, _det("hdfc"), 1.0))
    monkeypatch.setattr(po, "_get_bank_processor",
                        lambda k: _StubProcessor if k == "hdfc" else None)
    return po


def test_orchestrator_off_mode_block_and_legacy_path(monkeypatch, tmp_path):
    import app.services.pipeline_orchestrator as po
    monkeypatch.setattr(po, "_get_bank_processor",
                        lambda k: _StubProcessor if k == "hdfc" else None)
    monkeypatch.setattr(_StubProcessor, "process", _stub_process_ok)
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    out = po.process_statement(str(pdf), {"bank_name": "hdfc"}, output_dir=str(tmp_path))
    assert out["statement_engine"] == {"engine": "legacy", "mode": "off",
                                       "layout_id": None, "verdict": None,
                                       "needs_review": False, "review_codes": [],
                                       "anchored": False, "detected_bank": None,
                                       "advisory": None}
    assert out["status"] == "success"


def test_orchestrator_serves_v2_verified(monkeypatch, tmp_path):
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "hdfc")
    po = _patch_orchestrator(monkeypatch, _routed("VERIFIED"))
    monkeypatch.setattr(_StubProcessor, "process", _stub_process_ok)
    monkeypatch.setattr(_StubProcessor, "_run_pipeline", _stub_v2_pipeline)
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    out = po.process_statement(str(pdf), {"bank_name": "hdfc"}, output_dir=str(tmp_path))
    se = out["statement_engine"]
    assert se["engine"] == "v2" and se["verdict"] == "VERIFIED"
    assert se["anchored"] is True and se["layout_id"] == "hdfc_x"
    assert _V2_ROWS_SEEN and len(_V2_ROWS_SEEN[-1].transactions) == 3
    _V2_ROWS_SEEN.clear()


def test_orchestrator_review_serves_legacy_with_advisory(monkeypatch, tmp_path):
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "hdfc")
    po = _patch_orchestrator(monkeypatch, _routed("REVIEW"))
    monkeypatch.setattr(_StubProcessor, "process", _stub_process_ok)
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    out = po.process_statement(str(pdf), {"bank_name": "hdfc"}, output_dir=str(tmp_path))
    se = out["statement_engine"]
    assert se["engine"] == "legacy" and se["advisory"]["verdict"] == "REVIEW"
    assert out["status"] == "success"


def test_orchestrator_timeout_falls_back_to_legacy(monkeypatch, tmp_path):
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "hdfc")
    po = _patch_orchestrator(monkeypatch, None)
    monkeypatch.setattr(_StubProcessor, "process", _stub_process_ok)
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    out = po.process_statement(str(pdf), {"bank_name": "hdfc"}, output_dir=str(tmp_path))
    assert out["statement_engine"]["engine"] == "legacy"
    assert out["status"] == "success"


def test_orchestrator_unsupported_bank_without_serve_still_raises(monkeypatch, tmp_path):
    import app.services.pipeline_orchestrator as po
    monkeypatch.setattr(po, "_get_bank_processor", lambda k: None)
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    with pytest.raises(po.UnsupportedBankError):
        po.process_statement(str(pdf), {"bank_name": "madeup bank"}, output_dir=str(tmp_path))


def test_orchestrator_audit_event_written(monkeypatch, tmp_path):
    monkeypatch.setenv(shadow.MODE_ENV, "serve")
    monkeypatch.setenv(serve.SERVE_BANKS_ENV, "hdfc")
    po = _patch_orchestrator(monkeypatch, _routed("VERIFIED"))
    monkeypatch.setattr(_StubProcessor, "process", _stub_process_ok)
    monkeypatch.setattr(_StubProcessor, "_run_pipeline", _stub_v2_pipeline)
    events = []
    audit = SimpleNamespace(create_job_event=lambda *a, **k: events.append(k))
    pdf = tmp_path / "s.pdf"
    pdf.write_bytes(b"%PDF-fake")
    po.process_statement(str(pdf), {"bank_name": "hdfc"}, output_dir=str(tmp_path),
                         audit_service=audit, job_id="j1")
    assert events and events[0]["event_type"] == "v2_serve_decision"
    assert events[0]["metadata"]["engine"] == "v2"
    # block carries no statement content
    flat = str(events[0]["metadata"])
    for banned in ("description", "9999", "T0", "T1", "T2"):
        assert banned not in flat


# --------------------------------------------------------------------------- #
# Excel Statement Check sheet                                                  #
# --------------------------------------------------------------------------- #

def _tiny_workbook(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "x"
    ws["B2"] = 42
    wb.save(path)


def test_statement_check_sheet_appended_only_when_wanted(tmp_path):
    xlsx = tmp_path / "w.xlsx"
    _tiny_workbook(str(xlsx))
    block = serve.engine_block(engine="v2", mode="serve", routed=_routed("REVIEW"),
                               verdict="REVIEW", needs_review=True,
                               codes=["NO_INDEPENDENT_ANCHOR"])
    assert serve.wants_statement_check(block) is True
    assert serve.append_statement_check_sheet(str(xlsx), block) is True
    from openpyxl import load_workbook
    wb = load_workbook(str(xlsx))
    assert wb.sheetnames == ["Summary", "Statement Check"]
    assert wb["Summary"]["B2"].value == 42  # existing cells untouched


def test_statement_check_not_wanted_for_plain_legacy():
    block = serve.minimal_block("off")
    assert serve.wants_statement_check(block) is False


def test_workbook_byte_identical_when_off(tmp_path):
    xlsx = tmp_path / "w.xlsx"
    _tiny_workbook(str(xlsx))
    before = xlsx.read_bytes()
    block = serve.minimal_block("off")
    if serve.wants_statement_check(block):
        serve.append_statement_check_sheet(str(xlsx), block)
    assert xlsx.read_bytes() == before


def test_statement_check_lands_after_real_generator_sheets(tmp_path):
    """End-to-end on the real LiteExcelGenerator: synthetic rows -> 9 sheets,
    then the Statement Check sheet appended last without touching them."""
    from app.services.pipeline.reporting.lite_excel_generator import LiteExcelGenerator
    from openpyxl import load_workbook

    txns = [{
        "date": "2023-04-0%d" % i, "description": "TXN %d" % i,
        "debit": 10.0 * i if i % 2 else None, "credit": 5.0 * i if i % 2 else None,
        "balance": 100.0 + i, "category": "Others", "confidence": "",
        "recurring": "No", "ref_no": "R%d" % i,
    } for i in range(1, 6)]
    meta = {"bank_name": "Test Bank", "opening_balance": 100, "closing_balance": 105,
            "total_transactions": 5, "statement_from": "2023-04-01",
            "statement_to": "2023-04-30"}
    xlsx = str(tmp_path / "report.xlsx")
    LiteExcelGenerator().generate(txns, meta, xlsx)

    before = load_workbook(xlsx)
    names_before = list(before.sheetnames)
    hash_before = hash(tuple(
        (ws.title, tuple(tuple(c.value for c in row) for row in ws.iter_rows()))
        for ws in before.worksheets
    ))

    block = serve.engine_block(engine="v2", mode="serve", routed=_routed("VERIFIED"),
                               verdict="VERIFIED")
    assert serve.append_statement_check_sheet(xlsx, block) is True

    after = load_workbook(xlsx)
    assert after.sheetnames == names_before + ["Statement Check"]
    hash_after = hash(tuple(
        (ws.title, tuple(tuple(c.value for c in row) for row in ws.iter_rows()))
        for ws in after.worksheets[: len(names_before)]
    ))
    assert hash_after == hash_before  # the 9 original sheets are cell-identical


def test_statement_check_flags_needs_review_when_advisory_failed(tmp_path):
    """Legacy-served workbook next to a FAILED engine advisory must say it needs review (found in overnight verification:
    a FAILED statement showed 'Needs review: no')."""
    xlsx = tmp_path / "w.xlsx"
    _tiny_workbook(str(xlsx))
    adv = {"verdict": "FAILED", "review_codes": ["CHAIN_BREAK", "SIDE_MISMATCH"]}
    block = serve.engine_block(engine="legacy", mode="serve", routed=None, verdict=None, needs_review=False,
                               codes=[], advisory=adv)
    assert serve.append_statement_check_sheet(str(xlsx), block) is True
    from openpyxl import load_workbook
    rows = {r[0]: r[1] for r in load_workbook(str(xlsx))["Statement Check"].iter_rows(values_only=True) if r and r[0]}
    assert rows["Verdict"] == "FAILED" and rows["Needs review"] == "yes"
    assert "legacy parser" in rows["Note"]


def test_statement_check_advisory_review_does_not_force_needs_review(tmp_path):
    xlsx = tmp_path / "w.xlsx"
    _tiny_workbook(str(xlsx))
    block = serve.engine_block(engine="legacy", mode="serve", routed=None, verdict=None, needs_review=False,
                               codes=[], advisory={"verdict": "REVIEW", "review_codes": ["NO_INDEPENDENT_ANCHOR"]})
    serve.append_statement_check_sheet(str(xlsx), block)
    from openpyxl import load_workbook
    rows = {r[0]: r[1] for r in load_workbook(str(xlsx))["Statement Check"].iter_rows(values_only=True) if r and r[0]}
    assert rows["Needs review"] == "no" and "Note" not in rows


def test_row_count_guard_flags_workbook_that_differs_from_engine():
    from types import SimpleNamespace as NS
    from app.services.statement_engine import serve
    routed = NS(result=NS(transactions=[1] * 140))
    blk = {"engine": "v2", "verdict": "VERIFIED", "needs_review": False, "anchored": True, "review_codes": []}
    out = serve.apply_row_count_guard(dict(blk), routed, 142)
    assert out["needs_review"] is True and out["anchored"] is False
    assert "WORKBOOK_ROW_COUNT_MISMATCH" in out["review_codes"]


def test_row_count_guard_leaves_matching_or_legacy_blocks_alone():
    from types import SimpleNamespace as NS
    from app.services.statement_engine import serve
    routed = NS(result=NS(transactions=[1] * 140))
    ok = {"engine": "v2", "needs_review": False, "anchored": True, "review_codes": []}
    assert serve.apply_row_count_guard(dict(ok), routed, 140) == ok
    leg = {"engine": "legacy", "needs_review": False}
    assert serve.apply_row_count_guard(dict(leg), routed, 1) == leg
