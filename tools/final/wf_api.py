"""N9 workflow API leg on rc (:8601): workbook checks, two-user isolation,
retention text, mid-job kill, log scan. Synthetic files only; no content printed.
usage: python tools/final/wf_api.py
"""
import io, json, os, subprocess, sys, time
import requests

BE = "http://127.0.0.1:8601"
KC = "http://localhost:8080"
KIT = "tools/final/demo_pdfs"
OUT = "tools/final/.local"
os.makedirs(OUT, exist_ok=True)
ENV = {}
for _p in ("X:/Airco Projects/Airco Insights fintech/local.env",
           "X:/Airco Projects/Airco Insights fintech/.env"):
    if os.path.exists(_p):
        for line in open(_p, encoding="utf-8"):
            m = line.strip().split("=", 1)
            if len(m) == 2:
                ENV.setdefault(m[0], m[1].strip('"'))
USER, PW = ENV.get("TEST_USER_EMAIL") or os.environ.get("TEST_USER_EMAIL", "test@airco.com"), \
           ENV.get("TEST_USER_PASSWORD") or os.environ.get("TEST_USER_PASSWORD", "")
SECRET = ENV.get("KEYCLOAK_CLIENT_SECRET") or os.environ.get("KEYCLOAK_CLIENT_SECRET", "")
R = {}

def login(u=USER, p=PW):
    r = requests.post(f"{KC}/realms/airco-insights/protocol/openid-connect/token",
                      data={"grant_type": "password", "client_id": "frontend-app",
                            "client_secret": SECRET, "username": u, "password": p, "scope": "openid"},
                      timeout=30)
    return {"Authorization": "Bearer " + r.json()["access_token"]}

def run_flow(H, path, bank_override=None, password=None, wait=420):
    files = {"files": (os.path.basename(path), open(path, "rb").read(), "application/pdf")}
    r = requests.post(f"{BE}/api/statements/analyze", headers=H, files=files, timeout=120)
    if r.status_code == 401:
        H.update(login()); files = {"files": (os.path.basename(path), open(path, "rb").read(), "application/pdf")}
        r = requests.post(f"{BE}/api/statements/analyze", headers=H, files=files, timeout=120)
    a = r.json()
    doc = (a.get("documents") or [{}])[0]
    out = {"analyze_http": r.status_code, "state": doc.get("state"), "detail": doc.get("detail_state"),
           "bank": (doc.get("bank") or {}).get("id"), "workflow": doc.get("workflow_status")}
    aid = a.get("analysis_id")
    did = doc.get("document_id")
    if not aid or not did or doc.get("workflow_status") != "AWAITING_CONFIRMATION":
        return out
    doc_payload = {"document_id": did, "action": "CONFIRM"}
    if bank_override:
        doc_payload = {"document_id": did, "action": "OVERRIDE", "bank_id": bank_override}
    if password:
        doc_payload["password"] = password
    rc = requests.post(f"{BE}/api/statements/{aid}/confirm", headers=H, json={"documents": [doc_payload]}, timeout=90)
    out["confirm_http"] = rc.status_code
    try:
        cj = rc.json()
        res = cj.get("results") or cj
        out["confirm_ok"] = any(x.get("ok") for x in res) if isinstance(res, list) else True
    except Exception:
        out["confirm_ok"] = False
        out["confirm_body"] = rc.text[:140]
        return out
    if not out["confirm_ok"]:
        out["confirm_body"] = str(cj)[:160]
        return out
    rp = requests.post(f"{BE}/api/statements/{aid}/process", headers=H,
                       json={"document_ids": [did], "full_name": "N9 Verify"}, timeout=90)
    out["process_http"] = rp.status_code
    try:
        pj = rp.json()
    except Exception:
        out["process_body"] = rp.text[:140]; return out
    started = pj.get("started") or []
    if not started:
        out["skipped"] = pj.get("skipped"); return out
    job = started[0]["job_id"]
    out["job_id"] = job
    t0 = time.time()
    while time.time() - t0 < wait:
        j = requests.get(f"{BE}/api/v1/jobs/{job}", headers=H, timeout=30).json()
        if j.get("status") in ("completed", "failed"):
            break
        time.sleep(6)
    em = ((j.get("result_data") or {}).get("statement_engine")) or {}
    out.update({"job_status": j.get("status"), "engine": em.get("engine"), "verdict": em.get("verdict"),
                "needs_review": em.get("needs_review"), "review_codes": em.get("review_codes"),
                "txns": (j.get("result_data") or {}).get("transaction_count")})
    d = requests.get(f"{BE}/api/v1/jobs/{job}/download", headers=H, timeout=60)
    out["dl_http"] = d.status_code
    if d.status_code == 200:
        out["wb"] = workbook_check(d.content)
    return out

def workbook_check(buf):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(buf), data_only=False)
    errs, tx_rows, sc = [], 0, {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and any(
                        e in c.value for e in ("REF", "VALUE", "DIV", "NAME", "N/A", "NULL", "NUM")):
                    errs.append(f"{ws.title}!{c.coordinate}")
    if "Transactions" in wb.sheetnames:
        ws = wb["Transactions"]
        tx_rows = max(0, ws.max_row - 1)
    if "Statement Check" in wb.sheetnames:
        ws = wb["Statement Check"]
        for row in ws.iter_rows(max_col=2):
            k = str(row[0].value or "")
            if k:
                sc[k[:40]] = str(row[1].value)[:60]
    return {"sheets": wb.sheetnames, "tx_rows": tx_rows, "formula_errors": errs[:8], "check": sc}

H = login()

# --- 6: VERIFIED workbook (canara/karnataka kit files produced VERIFIED locally) ---
for b in ("canara", "karnataka", "au_sfb"):
    r = run_flow(H, f"{KIT}/demo_{b}.pdf")
    if r.get("verdict") == "VERIFIED":
        R["s6_verified_wb"] = r; break
    R.setdefault("s6_attempts", []).append((b, r.get("verdict"), r.get("job_status")))
if "s6_verified_wb" not in R:
    R["s6_verified_wb"] = "none of canara/karnataka/au_sfb produced VERIFIED on rc"

# --- 7: REVIEW workbook says why ---
for b in ("iob", "south_indian", "pnb", "hdfc"):
    r = run_flow(H, f"{KIT}/demo_{b}.pdf")
    if r.get("verdict") == "REVIEW":
        R["s7_review_wb"] = r; break
    R.setdefault("s7_attempts", []).append((b, r.get("verdict"), r.get("job_status")))
if "s7_review_wb" not in R:
    R["s7_review_wb"] = "no REVIEW job produced"

# --- 9: two-user isolation (create user2 via KC admin) ---
try:
    adm = requests.post(f"{KC}/realms/master/protocol/openid-connect/token",
                        data={"grant_type": "password", "client_id": "admin-cli",
                              "username": ENV.get("KEYCLOAK_ADMIN", "admin"),
                              "password": ENV.get("KEYCLOAK_ADMIN_PASSWORD", "")}, timeout=30)
    at = adm.json()["access_token"]
    AH = {"Authorization": "Bearer " + at, "Content-Type": "application/json"}
    requests.post(f"{KC}/admin/realms/airco-insights/users", headers=AH, json={
        "username": "n9test2@airco.com", "email": "n9test2@airco.com", "enabled": True,
        "firstName": "N9", "lastName": "Test2", "emailVerified": True, "requiredActions": [],
        "credentials": [{"type": "password", "value": "N9Test!234", "temporary": False}]}, timeout=30)
    # find user2 id to reset creds + ensure profile fields (KC blocks ROPC without them)
    u2 = requests.get(f"{KC}/admin/realms/airco-insights/users?username=n9test2@airco.com", headers=AH, timeout=30).json()
    if u2:
        requests.put(f"{KC}/admin/realms/airco-insights/users/{u2[0]['id']}", headers=AH,
                     json={"firstName": "N9", "lastName": "Test2", "emailVerified": True,
                           "requiredActions": []}, timeout=30)
        requests.put(f"{KC}/admin/realms/airco-insights/users/{u2[0]['id']}/reset-password", headers=AH,
                     json={"type": "password", "value": "N9Test!234", "temporary": False}, timeout=30)
    H2 = login("n9test2@airco.com", "N9Test!234")
    r2 = run_flow(H2, f"{KIT}/demo_canara.pdf")
    R["s9_user2_flow"] = {k: r2.get(k) for k in ("job_status", "verdict", "job_id")}
    j1 = requests.get(f"{BE}/api/jobs?limit=50", headers=H, timeout=30).json()
    j2 = requests.get(f"{BE}/api/jobs?limit=50", headers=H2, timeout=30).json()
    n1 = j1 if isinstance(j1, list) else j1.get("jobs", [])
    n2 = j2 if isinstance(j2, list) else j2.get("jobs", [])
    ids1 = {str(j.get("id")) for j in n1}
    ids2 = {str(j.get("id")) for j in n2}
    u2job = str(r2.get("job_id"))
    R["s9_isolation"] = {
        "u1_jobs": len(ids1), "u2_jobs": len(ids2),
        "overlap": len(ids1 & ids2),
        "u2_sees_only_own": bool(n2) and all(j.get("user_id") == n2[0].get("user_id") for j in n2),
        "u1_sees_only_own": len({j.get("user_id") for j in n1}) == 1,
    }
except Exception as e:
    R["s9_error"] = str(e)[:200]

# --- 10: retention / expiry messaging ---
try:
    txt = requests.get(f"{BE}/api/statements/banks", headers=H, timeout=30)
    # retention hints live in API responses / docs; check analysis response fields
    r = requests.post(f"{BE}/api/statements/analyze", headers=H,
                      files={"files": ("demo_canara.pdf", open(f"{KIT}/demo_canara.pdf", "rb").read(), "application/pdf")}, timeout=120)
    a = r.json(); doc = (a.get("documents") or [{}])[0]
    R["s10_analysis_fields"] = sorted(doc.keys())[:24]
    R["s10_retention_text"] = [v for v in json.dumps(a).split('"') if "day" in v.lower() or "retention" in v.lower() or "expir" in v.lower()][:6]
except Exception as e:
    R["s10_error"] = str(e)[:160]

# --- 15: worker killed mid-job -> honest state + retry ---
try:
    files = {"files": ("demo_indian_bank.pdf", open(f"{KIT}/demo_indian_bank.pdf", "rb").read(), "application/pdf")}
    r = requests.post(f"{BE}/api/statements/analyze", headers=H, files=files, timeout=120)
    a = r.json(); doc = (a.get("documents") or [{}])[0]
    aid = a.get("analysis_id"); did = doc.get("document_id")
    if aid and did:
        requests.post(f"{BE}/api/statements/{aid}/confirm", headers=H,
                      json={"documents": [{"document_id": did, "action": "CONFIRM"}]}, timeout=60)
        rp = requests.post(f"{BE}/api/statements/{aid}/process", headers=H,
                           json={"document_ids": [did], "full_name": "N9 Kill"}, timeout=60)
        job = ((rp.json() or {}).get("started") or [{}])[0].get("job_id")
        if job:
            time.sleep(10)  # let it enter processing
            st = requests.get(f"{BE}/api/v1/jobs/{job}", headers=H, timeout=30).json().get("status")
            subprocess.run(["docker", "restart", "rc_backend"], check=False, capture_output=True, timeout=90)
            time.sleep(35)
            for _ in range(40):
                j = requests.get(f"{BE}/api/v1/jobs/{job}", headers=H, timeout=30).json()
                if j.get("status") in ("completed", "failed"):
                    break
                time.sleep(8)
            R["s15_kill"] = {"status_at_kill": st, "final": j.get("status"),
                             "honest": j.get("status") in ("completed", "failed")}
except Exception as e:
    R["s15_error"] = str(e)[:160]

# --- 16: backend log scan for statement content / secrets ---
try:
    lg = subprocess.run(["docker", "logs", "rc_backend", "--tail", "4000"],
                        capture_output=True, timeout=60)
    blob = (lg.stdout or b"").decode("utf-8", "replace") + (lg.stderr or b"").decode("utf-8", "replace")
    import re
    pats = {
        "secret_words": len(re.findall(r"keycloak.*secret|client_secret|Test123|demo1234|guest:guest", blob, re.I)),
        "amount_like": len(re.findall(r"\b\d{1,3}(?:,\d{3})+\.\d{2}\b", blob)),
        "ifsc_like": len(re.findall(r"\b[A-Z]{4}0[A-Z0-9]{6}\b", blob)),
        "acct_like": len(re.findall(r"\b\d{10,16}\b", blob)),
        "narration_words": len(re.findall(r"UPI/TEST|Test Holder", blob)),
    }
    R["s16_log_scan"] = pats
except Exception as e:
    R["s16_error"] = str(e)[:160]

json.dump(R, open(f"{OUT}/wf_api_results.json", "w"), indent=1)
print(json.dumps({k: (str(v)[:180]) for k, v in R.items()}, indent=1)[:6000])
