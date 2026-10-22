# P2 OCR e2e through the REAL API: analyze -> confirm(OVERRIDE other) -> process -> download.
# No PII printed: sha8 + counts + codes only.
import os, sys, json, time, hashlib, urllib.request, urllib.error, uuid

BASE = os.environ.get("R5B_BACKEND", "http://localhost:8701")
KC = "http://localhost:8080"
REALM = "airco-insights"
CLIENT = "frontend-app"
USER = os.environ.get("R5B_USER", "devin-e2e")
PW = os.environ["R5B_PASSWORD"]
SECRET = os.environ["R5B_SECRET"]


def token():
    data = urllib.parse.urlencode({
        "grant_type": "password", "client_id": CLIENT, "client_secret": SECRET,
        "username": USER, "password": PW, "scope": "openid"}).encode()
    with urllib.request.urlopen(f"{KC}/realms/{REALM}/protocol/openid-connect/token", data=data) as r:
        return json.load(r)["access_token"]


def api(tk, method, path, body=None, raw=None, filename=None, ctype="application/pdf"):
    url = f"{BASE}{path}"
    headers = {"Authorization": f"Bearer {tk}"}
    if raw is not None:
        boundary = uuid.uuid4().hex
        pre = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"{filename}\"\r\n"
               f"Content-Type: {ctype}\r\n\r\n").encode()
        payload = pre + raw + f"\r\n--{boundary}--\r\n".encode()
        headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    elif body is not None:
        payload = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    else:
        payload = None
    req = urllib.request.Request(url, data=payload, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as r:
            b = r.read()
            return r.status, (json.loads(b) if b and b[:1] in (b"{", b"[") else b)
    except urllib.error.HTTPError as e:
        b = e.read()
        try:
            return e.code, json.loads(b)
        except Exception:
            return e.code, b[:200]


def analyze(tk, pdf_path):
    raw = open(pdf_path, "rb").read()
    return api(tk, "POST", "/api/statements/analyze",
               raw=raw, filename=os.path.basename(pdf_path))


def main():
    pdf = sys.argv[1]
    sha8 = hashlib.sha256(open(pdf, "rb").read()).hexdigest()[:8]
    tk = token()
    st, res = analyze(tk, pdf)
    print("analyze:", st)
    if st != 200:
        print(json.dumps(res)[:400]); return 1
    aid = res["analysis_id"]
    docs = res.get("documents") or []
    d = docs[0]
    print("doc:", d["document_id"][:8], "state:", d.get("state"), "capability:",
          (d.get("capability") or {}).get("capability") if isinstance(d.get("capability"), dict) else d.get("capability"))
    cap = d.get("capability") or {}
    print("capability:", cap.get("capability"), "pages:", cap.get("page_count"))
    bank = d.get("bank") or {}
    print("detected bank:", bank.get("id"), "confidence:", bank.get("confidence"), "supported:", bank.get("supported"))
    # confirm: OVERRIDE other (scanned -> T4)
    st, c = api(tk, "POST", f"/api/statements/{aid}/confirm",
                {"documents": [{"document_id": d["document_id"], "action": "OVERRIDE", "bank_id": "other"}]})
    print("confirm:", st, json.dumps(c.get("results"))[:160] if isinstance(c, dict) else c)
    # process
    st, p = api(tk, "POST", f"/api/statements/{aid}/process",
                {"document_ids": [d["document_id"]], "full_name": "Test User"})
    print("process:", st)
    started = (p or {}).get("started") or []
    if not started:
        print(json.dumps(p)[:400]); return 1
    job = started[0]["job_id"]
    print("job:", str(job)[:12])
    # poll
    final = None
    for _ in range(90):
        time.sleep(2)
        st, j = api(tk, "GET", f"/api/jobs/{job}")
        if st != 200:
            print("job poll:", st, json.dumps(j)[:200]); return 1
        status = j.get("status")
        if status in ("completed", "failed", "done", "error"):
            final = j
            break
    if final is None:
        print("job did not finish in 180s"); return 1
    print("job status:", final.get("status"))
    err = final.get("error") or {}
    if err:
        print("error code:", err.get("code"), "| stage:", err.get("stage"), "| msg:", str(err.get("error"))[:120])
    se = final.get("statement_engine") or {}
    if se:
        print("engine:", se.get("engine"), "verdict:", se.get("verdict"), "anchored:", se.get("anchored"),
              "needs_review:", se.get("needs_review"))
        print("review_codes:", se.get("review_codes"))
        ocr = se.get("ocr") or {}
        print("ocr block:", json.dumps({k: ocr.get(k) for k in ("pages_ok", "failed_pages", "partial", "mode")})[:200])
    stats = final.get("stats") or {}
    print("stats: tx=", stats.get("total_transactions"), "coverage=", stats.get("coverage_percent"))
    dl = final.get("download_url") or final.get("excel_url")
    print("download path present:", bool(dl))
    if dl:
        st2, wb = api(tk, "GET", dl if dl.startswith("/") else f"/api/jobs/{job}/download")
        print("download:", st2, "bytes:", len(wb) if isinstance(wb, bytes) else "?")
        if isinstance(wb, bytes) and st2 == 200:
            out = rf"C:/Users/admin/AppData/Local/Temp/r5b_wb_{sha8}.xlsx"
            open(out, "wb").write(wb)
            print("saved:", out)
            import openpyxl
            w = openpyxl.load_workbook(out)
            print("sheets:", len(w.sheetnames), w.sheetnames)
            sc = w["Statement Check"]
            cells = {r[0]: r[1] for r in sc.iter_rows(values_only=True) if r and r[0]}
            keep = {}
            for k, v in cells.items():
                sk = str(k)
                if any(t in sk.lower() for t in ("engine", "verdict", "ocr", "review", "pages", "independently", "incomplete", "transactions")):
                    keep[sk] = v
            print("Statement Check (filtered):", json.dumps(keep, default=str)[:600])
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
