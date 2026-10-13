"""verify2 workbook reconciliation: Transactions sheet vs independent audit anchors.
usage: python tools/verify2/wb_verify.py <dl_dir> <audit.jsonl> [<audit2.jsonl>...]"""
import json, sys, glob, re
from decimal import Decimal
import openpyxl

dl_dir, audits = sys.argv[1], {}
for f in sys.argv[2:]:
    for ln in open(f, encoding='utf-8'):
        r = json.loads(ln)
        audits[r['sha8']] = r

DATE = re.compile(r'\d{2}[-/ ][A-Za-z]{3}[-/ ]\d{4}|\d{2}[-/.]\d{2}[-/.]\d{2,4}|\d{4}-\d{2}-\d{2}')

def num(v):
    if v is None or v == '':
        return None
    s = str(v).replace(',', '').replace('₹', '').strip()
    try:
        return Decimal(s)
    except Exception:
        return None

for xp in sorted(glob.glob(dl_dir + '/*.xlsx')):
    s8 = xp.split('\\')[-1].split('/')[-1][:8]
    wb = openpyxl.load_workbook(xp)
    names = [n.lower() for n in wb.sheetnames]
    txname = next((n for n in wb.sheetnames if n.lower() == 'transactions'), None)
    rec = {'sha8': s8, 'sheets': len(wb.sheetnames)}
    if txname:
        ws = wb[txname]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        rows = [r for r in rows if r[0] is not None]
        rec['sheet_rows'] = len(rows)
        deb = sum(num(r[3]) or 0 for r in rows)
        cred = sum(num(r[4]) or 0 for r in rows)
        first, last = rows[0], rows[-1]
        imp_open = (num(last[5]) or 0)
        r0bal = num(first[5])
        if r0bal is not None:
            imp_open = r0bal - (num(first[4]) or 0) + (num(first[3]) or 0)
        rec['wb_implied_open'] = str(imp_open)
        rec['wb_close'] = str(num(last[5]))
        rec['wb_debits'] = str(deb)
        rec['wb_credits'] = str(cred)
        # chain
        brk = 0
        for a, b in zip(rows, rows[1:]):
            side = (num(b[4]) or 0) - (num(b[3]) or 0)
            ab, bb = num(a[5]), num(b[5])
            if ab is not None and bb is not None and abs(bb - ab - side) > Decimal('0.011'):
                brk += 1
        rec['wb_chain_breaks'] = brk
        au = audits.get(s8)
        if au:
            rec['audit_rows'] = au.get('rows')
            rec['match_rows'] = rec['sheet_rows'] == au.get('rows')
            po = [Decimal(v) for v in au.get('printed_open') or []]
            pc = [Decimal(v) for v in au.get('printed_close') or []]
            rec['open_anchor_ok'] = any(abs(v - imp_open) <= Decimal('0.011') for v in po) if po else None
            rec['close_anchor_ok'] = any(abs(v - (num(last[5]) or 0)) <= Decimal('0.011') for v in pc) if pc else None
    print(json.dumps(rec))
